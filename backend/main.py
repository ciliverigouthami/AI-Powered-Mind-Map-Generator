from pydantic import BaseModel
from fastapi import FastAPI, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from dotenv import load_dotenv
from google import genai
from io import BytesIO
import os
import sqlite3
import json
import shutil
from datetime import datetime
from pptx import Presentation


# ============================================================
# ENVIRONMENT / GEMINI
# ============================================================

load_dotenv()

DEMO_MODE = os.getenv(
    "DEMO_MODE",
    "false"
).lower() == "true"

api_key = os.getenv("GEMINI_API_KEY")

if not api_key:
    raise RuntimeError(
        "GEMINI_API_KEY is not set. "
        "Please add it to backend/.env"
    )

client = genai.Client(
    api_key=api_key
)

DATABASE_FILE = "mindmaps.db"


# ============================================================
# SAFE GEMINI CALL
# ============================================================

def safe_gemini_call(
    prompt,
    response_format=None,
    fallback=None
):

    # --------------------------------------------------------
    # Manual Demo Mode
    # --------------------------------------------------------

    if DEMO_MODE and fallback is not None:

        return {
            "success": True,
            "output": fallback,
            "demo": True
        }

    try:

        if response_format:

            interaction = client.interactions.create(
                model="gemini-3.6-flash",
                input=prompt,
                response_format=response_format
            )

        else:

            interaction = client.interactions.create(
                model="gemini-3.6-flash",
                input=prompt
            )

        return {
            "success": True,
            "output": interaction.output_text,
            "demo": False
        }

    except Exception as error:

        error_message = str(error).lower()

        # ----------------------------------------------------
        # Automatic fallback if Gemini quota/rate limit
        # ----------------------------------------------------

        if (
            "quota" in error_message
            or "rate limit" in error_message
            or "429" in error_message
            or "resource exhausted" in error_message
        ):

            if fallback is not None:

                return {
                    "success": True,
                    "output": fallback,
                    "demo": True
                }

            return {
                "success": False,
                "error": (
                    "Gemini AI usage limit reached. "
                    "Please try again later or use Demo Mode."
                )
            }

        # ----------------------------------------------------
        # Automatic fallback for temporary failures
        # ----------------------------------------------------

        if fallback is not None:

            return {
                "success": True,
                "output": fallback,
                "demo": True
            }

        return {
            "success": False,
            "error": (
                "Gemini AI is temporarily unavailable. "
                "Please try again."
            )
        }


# ============================================================
# DATABASE
# ============================================================

def init_database():

    connection = sqlite3.connect(
        DATABASE_FILE
    )

    connection.execute("""
        CREATE TABLE IF NOT EXISTS mindmaps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            content TEXT,
            nodes TEXT NOT NULL,
            edges TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    connection.execute("""
        CREATE TABLE IF NOT EXISTS study_packs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            content TEXT,
            study_pack TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    connection.commit()
    connection.close()


init_database()


# ============================================================
# MIND MAP MODELS
# ============================================================

class MindMapRequest(BaseModel):

    content: str
    detail_level: str = "balanced"


class MindMapNode(BaseModel):

    id: str
    label: str
    type: str = "concept"


class MindMapEdge(BaseModel):

    source: str
    target: str
    label: str | None = None


class MindMapResponse(BaseModel):

    title: str
    summary: str
    nodes: list[MindMapNode]
    edges: list[MindMapEdge]


# ============================================================
# STUDY PACK MODELS
# ============================================================

class StudyPackTopic(BaseModel):

    title: str
    subtopics: list[str] = []


class StudyPackFlashcard(BaseModel):

    question: str
    answer: str


class StudyPackQuizQuestion(BaseModel):

    question: str
    options: list[str]
    correct_answer: str
    explanation: str


class StudyPackResponse(BaseModel):

    chapter_title: str
    summary: str
    key_points: list[str]
    topics: list[StudyPackTopic]
    nodes: list[MindMapNode]
    edges: list[MindMapEdge]
    quiz: list[StudyPackQuizQuestion]
    flashcards: list[StudyPackFlashcard]


class StudyPackRequest(BaseModel):

    content: str
    quiz_count: int = 10
    flashcard_count: int = 10


class SaveStudyPackRequest(BaseModel):

    title: str
    content: str
    study_pack: dict


# ============================================================
# STUDY TEXT CLEANING / CHUNKING
# ============================================================

def clean_study_text(text: str) -> str:

    lines = []

    previous_blank = False

    for raw_line in text.replace(
        "\r",
        ""
    ).split("\n"):

        line = " ".join(
            raw_line.split()
        ).strip()

        if not line:

            if not previous_blank:
                lines.append("")

            previous_blank = True

            continue

        lines.append(line)

        previous_blank = False

    return "\n".join(lines).strip()


def chunk_study_text(
    text: str,
    max_chars: int = 18000
) -> list[str]:

    text = clean_study_text(text)

    if len(text) <= max_chars:
        return [text]

    chunks = []

    current = []

    current_len = 0

    for paragraph in text.split("\n"):

        if not paragraph.strip():
            continue

        extra = len(paragraph) + 1

        if (
            current
            and current_len + extra > max_chars
        ):

            chunks.append(
                "\n".join(current)
            )

            current = []

            current_len = 0

        current.append(paragraph)

        current_len += extra

    if current:

        chunks.append(
            "\n".join(current)
        )

    return chunks


# ============================================================
# DEMO STUDY PACK
# ============================================================

def demo_study_pack(
    title: str = "Chapter Study Pack"
) -> dict:

    return {

        "chapter_title": title,

        "summary": (
            "This Demo Mode study pack shows how the "
            "uploaded chapter can be organized into a "
            "concise summary, key concepts, a mind map, "
            "quiz questions and flashcards."
        ),

        "key_points": [

            "Focus on the main definitions and concepts.",

            "Review the relationships between major topics.",

            "Practice the concepts using questions and flashcards.",

            "Use the mind map for quick revision."

        ],

        "topics": [

            {
                "title": "Core Concepts",
                "subtopics": [
                    "Definitions",
                    "Principles"
                ]
            },

            {
                "title": "Applications",
                "subtopics": [
                    "Examples",
                    "Use cases"
                ]
            },

            {
                "title": "Revision",
                "subtopics": [
                    "Important points",
                    "Practice questions"
                ]
            }

        ],

        "nodes": [

            {
                "id": "root",
                "label": title,
                "type": "root"
            },

            {
                "id": "topic-1",
                "label": "Core Concepts",
                "type": "main"
            },

            {
                "id": "topic-2",
                "label": "Applications",
                "type": "main"
            },

            {
                "id": "topic-3",
                "label": "Revision",
                "type": "main"
            },

            {
                "id": "concept-1",
                "label": "Definitions",
                "type": "concept"
            },

            {
                "id": "concept-2",
                "label": "Principles",
                "type": "concept"
            }

        ],

        "edges": [

            {
                "source": "root",
                "target": "topic-1",
                "label": "covers"
            },

            {
                "source": "root",
                "target": "topic-2",
                "label": "includes"
            },

            {
                "source": "root",
                "target": "topic-3",
                "label": "supports"
            },

            {
                "source": "topic-1",
                "target": "concept-1",
                "label": "defines"
            },

            {
                "source": "topic-1",
                "target": "concept-2",
                "label": "explains"
            }

        ],

        "quiz": [

            {
                "question": (
                    "What is the best first step "
                    "when studying a chapter?"
                ),

                "options": [

                    "Identify the main concepts",

                    "Skip the headings",

                    "Ignore relationships",

                    "Memorize every sentence"

                ],

                "correct_answer":
                    "Identify the main concepts",

                "explanation": (
                    "Finding the main concepts creates "
                    "a useful structure for revision."
                )
            }

        ],

        "flashcards": [

            {
                "question":
                    "What should a mind map focus on?",

                "answer": (
                    "The main concepts, topics and meaningful "
                    "relationships in the study material."
                )
            }

        ]

    }


# ============================================================
# OTHER MODELS
# ============================================================

class SaveMindMapRequest(BaseModel):

    title: str
    content: str = ""
    nodes: list[dict]
    edges: list[dict]


class ExplainNodeRequest(BaseModel):

    node_label: str
    content: str = ""


class QuizRequest(BaseModel):

    content: str
    num_questions: int = 5


class AskAIRequest(BaseModel):

    content: str
    question: str


# ============================================================
# FASTAPI APP
# ============================================================

app = FastAPI(

    title="MindMap AI",

    description=(
        "AI-powered study material to interactive "
        "mind map generator"
    ),

    version="1.0.0"

)


# ============================================================
# CORS
# ============================================================

app.add_middleware(

    CORSMiddleware,

    allow_origins=[

        "http://localhost:5173",

        "http://127.0.0.1:5173"

    ],

    allow_credentials=True,

    allow_methods=["*"],

    allow_headers=["*"],

)


# ============================================================
# BASIC ENDPOINTS
# ============================================================

@app.get("/")
def root():

    return {

        "message":
            "MindMap AI backend is running",

        "status":
            "success",

        "demo_mode":
            DEMO_MODE

    }


@app.get("/health")
def health_check():

    return {

        "status":
            "healthy",

        "service":
            "MindMap AI Backend",

        "demo_mode":
            DEMO_MODE

    }


@app.get("/api/test")
def api_test():

    return {

        "message":
            "React and FastAPI communication is working!",

        "backend":
            "FastAPI",

        "frontend":
            "React"

    }


# ============================================================
# GEMINI TEST
# ============================================================

@app.get("/api/gemini-test")
def gemini_test():

    try:

        response = client.interactions.create(

            model="gemini-3.6-flash",

            input=(
                "Say hello to our AI Mind Map Generator "
                "in one short sentence."
            )

        )

        return {

            "message":
                response.output_text

        }

    except Exception as error:

        return {

            "error":
                str(error)

        }


# ============================================================
# SAVE MIND MAP
# ============================================================

@app.post("/api/save-mindmap")
def save_mindmap(
    request: SaveMindMapRequest
):

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        cursor = connection.cursor()

        cursor.execute(

            """
            INSERT INTO mindmaps
            (title, content, nodes, edges, created_at)
            VALUES (?, ?, ?, ?, ?)
            """,

            (

                request.title,

                request.content,

                json.dumps(
                    request.nodes
                ),

                json.dumps(
                    request.edges
                ),

                datetime.now().isoformat()

            )

        )

        mindmap_id = cursor.lastrowid

        connection.commit()

        connection.close()

        return {

            "success":
                True,

            "message":
                "Mind map saved successfully.",

            "mindmap_id":
                mindmap_id

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not save mind map: {str(error)}"

        }


# ============================================================
# GET SAVED MIND MAPS
# ============================================================

@app.get("/api/mindmaps")
def get_mindmaps():

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        connection.row_factory = sqlite3.Row

        cursor = connection.cursor()

        cursor.execute("""

            SELECT id, title, created_at

            FROM mindmaps

            ORDER BY id DESC

        """)

        rows = cursor.fetchall()

        connection.close()

        mindmaps = []

        for row in rows:

            mindmaps.append({

                "id":
                    row["id"],

                "title":
                    row["title"],

                "created_at":
                    row["created_at"]

            })

        return {

            "success":
                True,

            "mindmaps":
                mindmaps

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not load saved mind maps: {str(error)}"

        }


# ============================================================
# GET ONE MIND MAP
# ============================================================

@app.get("/api/mindmaps/{mindmap_id}")
def get_mindmap(
    mindmap_id: int
):

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        connection.row_factory = sqlite3.Row

        cursor = connection.cursor()

        cursor.execute(

            """
            SELECT id, title, content, nodes, edges, created_at

            FROM mindmaps

            WHERE id = ?

            """,

            (mindmap_id,)

        )

        row = cursor.fetchone()

        connection.close()

        if not row:

            return {

                "success":
                    False,

                "error":
                    "Mind map not found."

            }

        return {

            "success":
                True,

            "mindmap": {

                "id":
                    row["id"],

                "title":
                    row["title"],

                "content":
                    row["content"],

                "nodes":
                    json.loads(
                        row["nodes"]
                    ),

                "edges":
                    json.loads(
                        row["edges"]
                    ),

                "created_at":
                    row["created_at"]

            }

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not load mind map: {str(error)}"

        }


# ============================================================
# DELETE MIND MAP
# ============================================================

@app.delete("/api/mindmaps/{mindmap_id}")
def delete_mindmap(
    mindmap_id: int
):

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        cursor = connection.cursor()

        cursor.execute(

            """
            DELETE FROM mindmaps

            WHERE id = ?

            """,

            (mindmap_id,)

        )

        if cursor.rowcount == 0:

            connection.close()

            return {

                "success":
                    False,

                "error":
                    "Mind map not found."

            }

        connection.commit()

        connection.close()

        return {

            "success":
                True,

            "message":
                "Mind map deleted successfully."

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not delete mind map: {str(error)}"

        }


# ============================================================
# PDF TEXT EXTRACTION + OCR
# ============================================================

def extract_pdf_text(
    pdf_bytes: bytes
) -> str:

    """
    Extract text from PDF.

    Processing order:

    1. PyMuPDF text extraction
    2. pypdf text extraction
    3. Tesseract OCR for scanned/image PDFs

    This allows both normal PDFs and scanned PDFs
    to be processed.
    """

    extracted_text = ""

    # --------------------------------------------------------
    # METHOD 1: PyMuPDF
    # --------------------------------------------------------

    try:

        import fitz

        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

        pages = []

        for page in document:

            text = page.get_text(
                "text"
            )

            if text and text.strip():

                pages.append(
                    text
                )

        document.close()

        extracted_text = "\n".join(
            pages
        )

        if len(
            extracted_text.strip()
        ) >= 50:

            return extracted_text

    except Exception:

        pass

    # --------------------------------------------------------
    # METHOD 2: pypdf
    # --------------------------------------------------------

    try:

        from pypdf import PdfReader

        reader = PdfReader(
            BytesIO(pdf_bytes)
        )

        pages = []

        for page in reader.pages:

            try:

                text = page.extract_text()

            except Exception:

                text = ""

            if text and text.strip():

                pages.append(
                    text
                )

        extracted_text = "\n".join(
            pages
        )

        if len(
            extracted_text.strip()
        ) >= 50:

            return extracted_text

    except Exception:

        pass

    # --------------------------------------------------------
    # METHOD 3: TESSERACT OCR
    # --------------------------------------------------------

    try:

        import fitz
        import pytesseract
        from PIL import Image

        # Find Tesseract automatically
        tesseract_exe = shutil.which(
            "tesseract"
        )

        # Windows fallback path
        if not tesseract_exe:

            windows_path = (
                r"C:\Program Files\Tesseract-OCR"
                r"\tesseract.exe"
            )

            if os.path.exists(
                windows_path
            ):

                tesseract_exe = windows_path

        if not tesseract_exe:

            return extracted_text

        pytesseract.pytesseract.tesseract_cmd = (
            tesseract_exe
        )

        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

        ocr_pages = []

        for page_number, page in enumerate(
            document,
            start=1
        ):

            try:

                # Higher resolution improves OCR
                matrix = fitz.Matrix(
                    2.5,
                    2.5
                )

                pixmap = page.get_pixmap(
                    matrix=matrix,
                    alpha=False
                )

                image_bytes = pixmap.tobytes(
                    "png"
                )

                image = Image.open(
                    BytesIO(image_bytes)
                )

                text = pytesseract.image_to_string(

                    image,

                    lang="eng",

                    config="--oem 3 --psm 6"

                )

                if text and text.strip():

                    ocr_pages.append(

                        f"\n--- Page {page_number} ---\n"
                        f"{text.strip()}"

                    )

            except Exception:

                continue

        document.close()

        ocr_text = "\n".join(
            ocr_pages
        )

        if ocr_text.strip():

            return ocr_text

    except Exception:

        pass

    return extracted_text


# ============================================================
# PDF UPLOAD
# ============================================================

@app.post("/api/upload-pdf")
async def upload_pdf(
    file: UploadFile = File(...)
):

    if not file.filename:

        return {

            "error":
                "No PDF file selected."

        }

    if not file.filename.lower().endswith(
        ".pdf"
    ):

        return {

            "error":
                "Please upload a PDF file."

        }

    try:

        pdf_bytes = await file.read()

        if not pdf_bytes:

            return {

                "error":
                    "The PDF file is empty."

            }

        extracted_text = extract_pdf_text(
            pdf_bytes
        )

        if not extracted_text.strip():

            return {

                "error": (
                    "Could not extract text from this PDF. "
                    "The PDF may contain scanned images or "
                    "handwritten content that OCR could not read."
                )

            }

        return {

            "filename":
                file.filename,

            "text":
                extracted_text,

            "ocr_available":
                bool(
                    shutil.which("tesseract")
                    or os.path.exists(
                        r"C:\Program Files\Tesseract-OCR"
                        r"\tesseract.exe"
                    )
                )

        }

    except Exception as error:

        return {

            "error":
                f"PDF processing failed: {str(error)}"

        }


# ============================================================
# DOCX UPLOAD
# ============================================================

@app.post("/api/upload-docx")
async def upload_docx(
    file: UploadFile = File(...)
):

    from docx import Document

    if not file.filename:

        return {

            "error":
                "No DOCX file selected."

        }

    if not file.filename.lower().endswith(
        ".docx"
    ):

        return {

            "error":
                "Please upload a DOCX file."

        }

    try:

        docx_bytes = await file.read()

        document = Document(
            BytesIO(docx_bytes)
        )

        extracted_text = ""

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                extracted_text += (
                    text + "\n"
                )

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:

                        row_text.append(
                            cell_text
                        )

                if row_text:

                    extracted_text += (
                        " | ".join(row_text)
                        + "\n"
                    )

        if not extracted_text.strip():

            return {

                "error":
                    "Could not extract text from this DOCX file."

            }

        return {

            "filename":
                file.filename,

            "text":
                extracted_text

        }

    except Exception as error:

        return {

            "error":
                f"DOCX processing failed: {str(error)}"

        }


# ============================================================
# GENERATE MIND MAP
# ============================================================

@app.post("/api/generate-mindmap")
def generate_mindmap(
    request: MindMapRequest
):

    detail_level = (
        request.detail_level
        .lower()
    )

    if detail_level not in [
        "quick",
        "balanced",
        "detailed"
    ]:

        detail_level = "balanced"

    # --------------------------------------------------------
    # QUICK
    # --------------------------------------------------------

    if detail_level == "quick":

        detail_instruction = """

Create a concise mind map.

NODE COUNT REQUIREMENT:

- Generate between 5 and 10 total nodes.
- Never exceed 10 nodes.

Focus only on:

- The root topic
- The most important main topics
- Essential concepts

Remove minor details, repeated information
and unnecessary examples.

"""

    # --------------------------------------------------------
    # DETAILED
    # --------------------------------------------------------

    elif detail_level == "detailed":

        detail_instruction = """

Create a detailed educational mind map.

NODE COUNT REQUIREMENT:

- Generate between 20 and 30 total nodes.
- Never exceed 30 nodes.

Include:

- The root topic
- Major topics
- Important subtopics
- Useful supporting concepts
- Meaningful relationships
- Important definitions
- Processes
- Classifications
- Key examples when useful

Prioritize important information instead of
creating a node for every sentence.

"""

    # --------------------------------------------------------
    # BALANCED
    # --------------------------------------------------------

    else:

        detail_instruction = """

Create a balanced educational mind map.

NODE COUNT REQUIREMENT:

- Generate between 12 and 20 total nodes.
- Never exceed 20 nodes.

Include:

- The root topic
- Important main topics
- Useful subtopics
- Important supporting concepts

Keep the map informative, concise
and easy to revise.

"""

    chapter_instruction = """

IMPORTANT FOR LARGE STUDY MATERIAL:

The input may come from an entire chapter,
large PDF or DOCX document.

Do NOT create one node for every sentence.

Instead:

1. Identify the overall subject.

2. Identify the most important sections or topics.

3. Group related information together.

4. Create a clear hierarchy:

Root
↓
Main Topics
↓
Subtopics
↓
Key Concepts

5. Remove repetitive information.

6. Avoid duplicate or very similar concepts.

7. Ignore unnecessary examples unless they
   improve understanding.

8. Prioritize information that is useful
   for student revision.

9. Give higher importance to:

   - headings
   - definitions
   - principles
   - processes
   - classifications
   - important relationships

10. Keep node labels short.

11. Do not create nodes for individual sentences.

"""

    prompt = f"""

You are an expert educational knowledge-map designer.

Your task is to convert study material into
a clear, structured and useful mind map for students.

{detail_instruction}

{chapter_instruction}

GENERAL RULES:

1. Create exactly ONE root node.

2. The root node MUST have:

id = "root"
type = "root"

3. The total number of nodes MUST follow
   the selected NODE COUNT REQUIREMENT above.

4. The node count includes the root node.

5. Create only the most meaningful concepts
   needed to understand the study material.

6. Create important main branches when enough
   information is available.

7. Add useful child concepts under main branches.

8. Keep node labels short and easy to understand.

9. Do not create duplicate concepts.

10. Use ONLY information supported by the
    study material.

11. Do not invent facts.

12. Create meaningful parent-child relationships.

13. Every edge must connect two existing node IDs.

14. EVERY edge MUST have a short relationship label.

15. Edge labels must contain only 1 to 3 simple words.

16. Never leave an edge label empty.

17. Do not create isolated nodes.

18. Prefer a clean hierarchical structure.

19. The mind map should help a student
    revise the material quickly.

20. Avoid extremely long node labels.

21. Use simple student-friendly language.

22. If the study material contains more information
    than can fit within the node limit, select the
    MOST IMPORTANT concepts.

23. Do not exceed the requested maximum node count.

STUDY MATERIAL:

{request.content}

"""

    demo_mindmap = {

        "title":
            "Machine Learning",

        "summary": (
            "Machine Learning enables computers "
            "to learn patterns from data and make "
            "predictions or decisions."
        ),

        "nodes": [

            {
                "id":
                    "root",
                "label":
                    "Machine Learning",
                "type":
                    "root"
            },

            {
                "id":
                    "2",
                "label":
                    "Supervised Learning",
                "type":
                    "topic"
            },

            {
                "id":
                    "3",
                "label":
                    "Unsupervised Learning",
                "type":
                    "topic"
            },

            {
                "id":
                    "4",
                "label":
                    "Reinforcement Learning",
                "type":
                    "topic"
            },

            {
                "id":
                    "5",
                "label":
                    "Training Data",
                "type":
                    "topic"
            },

            {
                "id":
                    "6",
                "label":
                    "Prediction",
                "type":
                    "topic"
            }

        ],

        "edges": [

            {
                "source":
                    "root",
                "target":
                    "2",
                "label":
                    "includes"
            },

            {
                "source":
                    "root",
                "target":
                    "3",
                "label":
                    "includes"
            },

            {
                "source":
                    "root",
                "target":
                    "4",
                "label":
                    "includes"
            },

            {
                "source":
                    "root",
                "target":
                    "5",
                "label":
                    "requires"
            },

            {
                "source":
                    "5",
                "target":
                    "6",
                "label":
                    "enables"
            }

        ]

    }

    try:

        ai_result = safe_gemini_call(

            prompt,

            response_format={

                "type":
                    "text",

                "mime_type":
                    "application/json",

                "schema":
                    MindMapResponse.model_json_schema()

            },

            fallback=json.dumps(
                demo_mindmap
            )

        )

        if not ai_result["success"]:

            return {

                "success":
                    False,

                "error":
                    ai_result["error"]

            }

        result = (
            MindMapResponse
            .model_validate_json(
                ai_result["output"]
            )
        )

        return result

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Mind map generation failed: {str(error)}"

        }


# ============================================================
# GENERATE STUDY PACK
# ============================================================

@app.post("/api/generate-study-pack")
def generate_study_pack(
    request: StudyPackRequest
):

    try:

        content = clean_study_text(
            request.content
        )

        if not content:

            return {

                "success":
                    False,

                "error":
                    "Please provide chapter content first."

            }

        quiz_count = max(
            5,
            min(
                request.quiz_count,
                15
            )
        )

        flashcard_count = max(
            5,
            min(
                request.flashcard_count,
                20
            )
        )

        chunks = chunk_study_text(
            content
        )

        # ----------------------------------------------------
        # LARGE DOCUMENT PROCESSING
        # ----------------------------------------------------

        if len(chunks) > 1:

            chunk_material = []

            for index, chunk in enumerate(
                chunks,
                start=1
            ):

                chunk_prompt = f"""

You are analyzing part {index}
of a student's chapter.

Extract only the most important study information
from this section.

Return concise plain text with:

- headings
- definitions
- key concepts
- processes
- classifications
- important relationships

Do not invent facts.

CHAPTER SECTION:

{chunk}

"""

                result = safe_gemini_call(

                    chunk_prompt,

                    fallback=chunk[:12000]

                )

                if result["success"]:

                    chunk_material.append(
                        result["output"]
                    )

            source = (
                "\n\n--- CHAPTER SECTION SUMMARY ---\n"
                .join(chunk_material)
            )

        else:

            source = content

        if len(source) > 120000:

            source = source[:120000]

        title_guess = (
            "Chapter Study Pack"
        )

        fallback = demo_study_pack(
            title_guess
        )

        fallback["quiz"] = (
            fallback["quiz"]
            * quiz_count
        )

        fallback["quiz"] = (
            fallback["quiz"][:quiz_count]
        )

        fallback["flashcards"] = (
            fallback["flashcards"]
            * flashcard_count
        )

        fallback["flashcards"] = (
            fallback["flashcards"][:flashcard_count]
        )

        prompt = f"""

You are an expert educational study-pack designer.

Analyze the provided complete chapter material
and create a useful study pack for a college student.

Return ONLY valid JSON matching the required schema.

Required output:

- chapter_title: concise chapter title

- summary: 1-2 concise paragraphs

- key_points: 6-12 high-value revision points

- topics: 5-10 major topics, each with
  1-5 short subtopics

- nodes: 12-20 total mind-map nodes

- exactly one root with id 'root'
  and type 'root'

- main topics should use type 'main'

- other concepts should use type 'concept'

- edges: meaningful parent-child relationships

- every edge must connect existing nodes

- every edge must have a 1-3 word label

- quiz: exactly {quiz_count} MCQs

- each quiz question must have
  4 options

- one correct answer

- concise explanations

- flashcards: exactly {flashcard_count}
  question-answer cards

Rules:

- Use ONLY information supported by
  the chapter material.

- Do not invent facts.

- Remove repeated information.

- Keep labels concise.

- Prioritize definitions, principles,
  processes, classifications, formulas,
  concepts and important relationships.

- Make the material useful for revision and exams.

CHAPTER MATERIAL:

{source}

"""

        ai_result = safe_gemini_call(

            prompt,

            response_format={

                "type":
                    "text",

                "mime_type":
                    "application/json",

                "schema":
                    StudyPackResponse.model_json_schema()

            },

            fallback=json.dumps(
                fallback
            )

        )

        if not ai_result["success"]:

            return {

                "success":
                    False,

                "error":
                    ai_result["error"]

            }

        result = (
            StudyPackResponse
            .model_validate_json(
                ai_result["output"]
            )
        )

        return {

            "success":
                True,

            "study_pack":
                result.model_dump(),

            "demo":
                ai_result.get(
                    "demo",
                    False
                )

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Study pack generation failed: {str(error)}"

        }


# ============================================================
# STUDY PACK TOPIC MAP
# ============================================================

@app.post("/api/study-pack-topic-map")
def study_pack_topic_map(
    request: AskAIRequest
):

    try:

        if (
            not request.content.strip()
            or not request.question.strip()
        ):

            return {

                "success":
                    False,

                "error":
                    "Chapter content and topic are required."

            }

        prompt = f"""

Create a focused mind map for the selected topic
from the chapter.

Use ONLY the provided chapter content.

Return JSON matching this schema.

Create exactly one root node.

Create 5-12 nodes total.

Use concise labels.

Create labeled relationships.

TOPIC:

{request.question}

CHAPTER CONTENT:

{request.content}

"""

        fallback = {

            "title":
                request.question,

            "summary":
                f"Focused study map for {request.question}.",

            "nodes": [

                {
                    "id":
                        "root",
                    "label":
                        request.question,
                    "type":
                        "root"
                },

                {
                    "id":
                        "1",
                    "label":
                        "Key Concepts",
                    "type":
                        "main"
                },

                {
                    "id":
                        "2",
                    "label":
                        "Applications",
                    "type":
                        "main"
                }

            ],

            "edges": [

                {
                    "source":
                        "root",
                    "target":
                        "1",
                    "label":
                        "covers"
                },

                {
                    "source":
                        "root",
                    "target":
                        "2",
                    "label":
                        "includes"
                }

            ]

        }

        class TopicMapResponse(BaseModel):

            title: str
            summary: str
            nodes: list[MindMapNode]
            edges: list[MindMapEdge]

        result = safe_gemini_call(

            prompt,

            response_format={

                "type":
                    "text",

                "mime_type":
                    "application/json",

                "schema":
                    TopicMapResponse.model_json_schema()

            },

            fallback=json.dumps(
                fallback
            )

        )

        if not result["success"]:

            return {

                "success":
                    False,

                "error":
                    result["error"]

            }

        parsed = (
            TopicMapResponse
            .model_validate_json(
                result["output"]
            )
        )

        return {

            "success":
                True,

            "map":
                parsed.model_dump(),

            "demo":
                result.get(
                    "demo",
                    False
                )

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Topic map generation failed: {str(error)}"

        }


# ============================================================
# SAVE STUDY PACK
# ============================================================

@app.post("/api/save-study-pack")
def save_study_pack(
    request: SaveStudyPackRequest
):

    try:

        content = clean_study_text(
            request.content
        )

        if not content:

            return {

                "success":
                    False,

                "error":
                    "No study pack content provided."

            }

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        connection.execute("""

            CREATE TABLE IF NOT EXISTS study_packs (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                title TEXT NOT NULL,

                content TEXT NOT NULL,

                study_pack TEXT NOT NULL,

                created_at TEXT NOT NULL

            )

        """)

        cursor = connection.cursor()

        cursor.execute(

            """

            INSERT INTO study_packs
            (title, content, study_pack, created_at)

            VALUES (?, ?, ?, ?)

            """,

            (

                request.title,

                content,

                json.dumps(
                    request.study_pack
                ),

                datetime.now().isoformat()

            )

        )

        pack_id = cursor.lastrowid

        connection.commit()

        connection.close()

        return {

            "success":
                True,

            "study_pack_id":
                pack_id

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not save study pack: {str(error)}"

        }


# ============================================================
# GET SAVED STUDY PACKS
# ============================================================

@app.get("/api/study-packs")
def get_study_packs():

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        connection.row_factory = sqlite3.Row

        connection.execute("""

            CREATE TABLE IF NOT EXISTS study_packs (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                title TEXT NOT NULL,

                content TEXT,

                study_pack TEXT NOT NULL,

                created_at TEXT NOT NULL

            )

        """)

        cursor = connection.cursor()

        cursor.execute("""

            SELECT id, title, created_at

            FROM study_packs

            ORDER BY id DESC

        """)

        rows = cursor.fetchall()

        connection.close()

        study_packs = []

        for row in rows:

            study_packs.append({

                "id":
                    row["id"],

                "title":
                    row["title"],

                "created_at":
                    row["created_at"]

            })

        return {

            "success":
                True,

            "study_packs":
                study_packs

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not load saved study packs: {str(error)}"

        }


# ============================================================
# GET ONE STUDY PACK
# ============================================================

@app.get("/api/study-packs/{study_pack_id}")
def get_study_pack(
    study_pack_id: int
):

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        connection.row_factory = sqlite3.Row

        connection.execute("""

            CREATE TABLE IF NOT EXISTS study_packs (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                title TEXT NOT NULL,

                content TEXT,

                study_pack TEXT NOT NULL,

                created_at TEXT NOT NULL

            )

        """)

        cursor = connection.cursor()

        cursor.execute(

            """

            SELECT
                id,
                title,
                content,
                study_pack,
                created_at

            FROM study_packs

            WHERE id = ?

            """,

            (study_pack_id,)

        )

        row = cursor.fetchone()

        connection.close()

        if not row:

            return {

                "success":
                    False,

                "error":
                    "Study Pack not found."

            }

        try:

            study_pack_data = json.loads(
                row["study_pack"]
            )

        except Exception:

            study_pack_data = {}

        return {

            "success":
                True,

            "study_pack": {

                "id":
                    row["id"],

                "title":
                    row["title"],

                "content":
                    row["content"] or "",

                "data":
                    study_pack_data,

                "created_at":
                    row["created_at"]

            }

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not load Study Pack: {str(error)}"

        }


# ============================================================
# DELETE STUDY PACK
# ============================================================

@app.delete("/api/study-packs/{study_pack_id}")
def delete_study_pack(
    study_pack_id: int
):

    try:

        connection = sqlite3.connect(
            DATABASE_FILE
        )

        cursor = connection.cursor()

        cursor.execute(

            """

            DELETE FROM study_packs

            WHERE id = ?

            """,

            (study_pack_id,)

        )

        if cursor.rowcount == 0:

            connection.close()

            return {

                "success":
                    False,

                "error":
                    "Study Pack not found."

            }

        connection.commit()

        connection.close()

        return {

            "success":
                True,

            "message":
                "Study Pack deleted successfully."

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not delete Study Pack: {str(error)}"

        }


# ============================================================
# EXPLAIN NODE
# ============================================================

@app.post("/api/explain-node")
def explain_node(
    request: ExplainNodeRequest
):

    try:

        prompt = f"""

You are an AI study assistant.

Explain the following study concept
in a simple and student-friendly way.

Concept:

{request.node_label}

Study material:

{request.content}

Return the answer in this exact JSON structure:

{{
"concept": "concept name",
"explanation": "simple explanation in 3-5 sentences",
"key_points": [
"key point 1",
"key point 2",
"key point 3"
],
"example": "simple example",
"real_world_application": "one real-world application"
}}

Rules:

- Use simple English.
- Keep the explanation concise.
- Do not use markdown.
- Do not add information unrelated to the concept.
- Return valid JSON only.

"""

        demo_explanation = {

            "concept":
                request.node_label,

            "explanation": (

                f"{request.node_label} is an important "
                "concept in the study material. It helps "
                "us understand the topic and how related "
                "ideas connect."

            ),

            "key_points": [

                (
                    f"Understanding {request.node_label} "
                    "helps build a strong foundation."
                ),

                (
                    f"{request.node_label} is connected "
                    "to other concepts in the topic."
                ),

                (
                    "It can be learned through examples "
                    "and practical applications."
                )

            ],

            "example": (

                f"A simple example of {request.node_label} "
                "can be understood by relating it to "
                "a real-world situation."

            ),

            "real_world_application": (

                f"{request.node_label} can be used to solve "
                "practical problems and understand "
                "real-world systems."

            )

        }

        ai_result = safe_gemini_call(

            prompt,

            fallback=json.dumps(
                demo_explanation
            )

        )

        if not ai_result["success"]:

            return {

                "success":
                    False,

                "error":
                    ai_result["error"]

            }

        result = json.loads(
            ai_result["output"]
        )

        return {

            "success":
                True,

            "explanation":
                result

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not explain the node: {str(error)}"

        }


# ============================================================
# GENERATE QUIZ
# ============================================================

@app.post("/api/generate-quiz")
def generate_quiz(
    request: QuizRequest
):

    try:

        question_count = max(

            3,

            min(
                request.num_questions,
                10
            )

        )

        prompt = f"""

You are an AI study assistant.

Create a multiple-choice quiz from
the following study material.

Study material:

{request.content}

Generate exactly {question_count} questions.

Return ONLY valid JSON in this exact structure:

{{
"questions": [
{{
"question": "Question text",
"options": [
"Option A",
"Option B",
"Option C",
"Option D"
],
"correct_answer": "Option A",
"explanation": "Short explanation of why this answer is correct."
}}
]
}}

Rules:

- Every question must have exactly 4 options.
- Only one option must be correct.
- Questions must be based on the provided study material.
- Use simple student-friendly English.
- Avoid duplicate questions.
- Do not use markdown.
- Return valid JSON only.

"""

        demo_questions = [

            {

                "question":
                    "What is the main purpose of Machine Learning?",

                "options": [

                    "To learn patterns from data",

                    "To replace all computers",

                    "To create hardware",

                    "To remove data"

                ],

                "correct_answer":
                    "To learn patterns from data",

                "explanation":
                    "Machine Learning allows computers "
                    "to learn patterns from data."

            },

            {

                "question":
                    "Which type of learning uses labelled data?",

                "options": [

                    "Supervised Learning",

                    "Unsupervised Learning",

                    "Reinforcement Learning",

                    "Random Learning"

                ],

                "correct_answer":
                    "Supervised Learning",

                "explanation":
                    "Supervised Learning uses labelled "
                    "training data."

            },

            {

                "question":
                    "Which learning approach works with rewards and penalties?",

                "options": [

                    "Supervised Learning",

                    "Unsupervised Learning",

                    "Reinforcement Learning",

                    "Manual Learning"

                ],

                "correct_answer":
                    "Reinforcement Learning",

                "explanation":
                    "Reinforcement Learning uses rewards "
                    "and penalties to guide learning."

            },

            {

                "question":
                    "What is training data used for?",

                "options": [

                    "Teaching a model",

                    "Turning off a computer",

                    "Deleting software",

                    "Formatting a disk"

                ],

                "correct_answer":
                    "Teaching a model",

                "explanation":
                    "Training data is used to help a "
                    "machine learning model learn patterns."

            },

            {

                "question":
                    "What can a trained machine learning model make?",

                "options": [

                    "Predictions",

                    "Only folders",

                    "Only images",

                    "Only passwords"

                ],

                "correct_answer":
                    "Predictions",

                "explanation":
                    "A trained model can use learned "
                    "patterns to make predictions."

            }

        ]

        demo_quiz = {

            "questions":
                demo_questions[:question_count]

        }

        ai_result = safe_gemini_call(

            prompt,

            fallback=json.dumps(
                demo_quiz
            )

        )

        if not ai_result["success"]:

            return {

                "success":
                    False,

                "error":
                    ai_result["error"]

            }

        result = json.loads(
            ai_result["output"]
        )

        return {

            "success":
                True,

            "questions":
                result.get(
                    "questions",
                    []
                )

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not generate quiz: {str(error)}"

        }


# ============================================================
# ASK AI
# ============================================================

@app.post("/api/ask-ai")
def ask_ai(
    request: AskAIRequest
):

    try:

        if not request.content.strip():

            return {

                "success":
                    False,

                "error":
                    "Please provide study material first."

            }

        if not request.question.strip():

            return {

                "success":
                    False,

                "error":
                    "Please enter a question."

            }

        prompt = f"""

You are an AI study assistant.

Answer the student's question using ONLY
the provided study material.

Study material:

{request.content}

Student question:

{request.question}

Rules:

- Answer based only on the study material.
- If the answer is not available in the study material,
  clearly say that the information is not present
  in the provided study material.
- Use simple, student-friendly English.
- Give a clear and concise answer.
- Do not use markdown.
- Do not invent facts.

"""

        demo_answer = (

            "This is a Demo Mode response. "

            "The AI would normally answer this question "

            "using the uploaded study material. "

            "During a presentation, the fallback keeps "

            "the application working even if Gemini is unavailable."

        )

        ai_result = safe_gemini_call(

            prompt,

            fallback=demo_answer

        )

        if not ai_result["success"]:

            return {

                "success":
                    False,

                "error":
                    ai_result["error"]

            }

        answer = (
            ai_result["output"]
            .strip()
        )

        return {

            "success":
                True,

            "answer":
                answer

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not answer the question: {str(error)}"

        }


# ============================================================
# STUDY TOOLS MODEL
# ============================================================

class StudyToolRequest(BaseModel):

    action: str

    content: str = ""

    topics: list[str] = []

    days: int = 7

    daily_minutes: int = 45

    difficulty: str = "medium"

    format: str = "exam"

    topic: str = ""

    weak_topics: list[str] = []

    question: str = ""

    language: str = "English"


# ============================================================
# STUDY TOOLS FALLBACK
# ============================================================

def study_tool_fallback(
    action: str,
    request: StudyToolRequest
):

    topics = request.topics or [

        "Core concepts",

        "Definitions",

        "Examples",

        "Applications"

    ]

    # --------------------------------------------------------
    # PLANNER
    # --------------------------------------------------------

    if action == "planner":

        return (

            f"{request.days}-day study plan\n"

            "Day 1: Learn the core concepts\n"

            "Day 2: Review definitions and key terms\n"

            "Day 3: Study relationships and examples\n"

            "Day 4: Practise questions\n"

            "Day 5: Revise weak areas\n"

            "Day 6: Take a self-test\n"

            "Day 7: Final revision and recall practice"

        )

    # --------------------------------------------------------
    # NOTES
    # --------------------------------------------------------

    if action == "notes":

        return (

            "EXAM-READY NOTES\n\n"

            "Key concepts\n- "

            + "\n- ".join(
                topics[:8]
            )

            + "\n\nUse the original material above "
              "to expand each concept with definitions, "
              "examples and applications."

        )

    # --------------------------------------------------------
    # REVISION
    # --------------------------------------------------------

    if action == "revision":

        focus = (

            request.topic

            or (

                request.weak_topics[0]

                if request.weak_topics

                else topics[0]

            )

        )

        return (

            f"Revision session: {focus}\n"

            "1. Recall the definition.\n"

            "2. Explain it without notes.\n"

            "3. Review one example.\n"

            "4. Solve two practice questions.\n"

            "5. Write a three-line summary.\n"

            "6. Revisit the topic tomorrow."

        )

    # --------------------------------------------------------
    # AI TUTOR
    # --------------------------------------------------------

    if action == "tutor":

        return (

            "Demo tutor response: review the supplied "
            "study material, identify the relevant concept, "
            "explain it in simple language, then verify "
            "your understanding with an example."

        )

    # --------------------------------------------------------
    # MULTILINGUAL
    # --------------------------------------------------------

    if action == "multilingual":

        return (

            f"Multilingual demo: the study notes can be "
            f"converted into {request.language}. "

            "Generate again with Gemini for a full faithful "
            "translation of the supplied material."

        )

    # --------------------------------------------------------
    # DIFFICULTY QUIZ
    # --------------------------------------------------------

    if action == "quiz":

        return [

            {

                "question":
                    "Which approach is best for learning the supplied material?",

                "options": [

                    "Active recall and practice",

                    "Reading once and stopping",

                    "Ignoring difficult topics",

                    "Memorizing without understanding"

                ],

                "correct_answer":
                    "Active recall and practice",

                "explanation":
                    "Active recall and practice strengthen "
                    "retrieval and understanding."

            }

        ]

    return "Study tool ready."


# ============================================================
# MULTIPLE FILE UPLOAD
# ============================================================

@app.post("/api/upload-multiple")
async def upload_multiple(
    files: list[UploadFile] = File(...)
):

    try:

        combined = []

        allowed = {

            "pdf",
            "docx",
            "pptx",
            "txt",
            "md"

        }

        for file in files[:10]:

            name = (
                file.filename
                or "file"
            )

            ext = (

                name.lower()
                .rsplit(".", 1)[-1]

                if "." in name

                else ""

            )

            if ext not in allowed:

                continue

            data = await file.read()

            # ------------------------------------------------
            # TXT / MARKDOWN
            # ------------------------------------------------

            if ext in {
                "txt",
                "md"
            }:

                text = data.decode(
                    "utf-8",
                    errors="ignore"
                )

            # ------------------------------------------------
            # PDF
            # ------------------------------------------------

            elif ext == "pdf":

                text = extract_pdf_text(
                    data
                )

            # ------------------------------------------------
            # DOCX
            # ------------------------------------------------

            elif ext == "docx":

                from docx import Document

                doc = Document(
                    BytesIO(data)
                )

                parts = []

                for paragraph in doc.paragraphs:

                    if paragraph.text.strip():

                        parts.append(
                            paragraph.text.strip()
                        )

                for table in doc.tables:

                    for row in table.rows:

                        cells = []

                        for cell in row.cells:

                            cell_text = (
                                cell.text.strip()
                            )

                            if cell_text:

                                cells.append(
                                    cell_text
                                )

                        if cells:

                            parts.append(
                                " | ".join(cells)
                            )

                text = "\n".join(
                    parts
                )

            # ------------------------------------------------
            # PPTX
            # ------------------------------------------------

            else:

                prs = Presentation(
                    BytesIO(data)
                )

                parts = []

                for slide in prs.slides:

                    for shape in slide.shapes:

                        if (

                            hasattr(
                                shape,
                                "text"
                            )

                            and shape.text.strip()

                        ):

                            parts.append(
                                shape.text.strip()
                            )

                text = "\n".join(
                    parts
                )

            if text.strip():

                combined.append(

                    f"\n===== {name} =====\n"
                    f"{text.strip()}"

                )

        if not combined:

            return {

                "success":
                    False,

                "error": (
                    "No readable text was found "
                    "in the uploaded files."
                )

            }

        return {

            "success":
                True,

            "text":
                "\n".join(
                    combined
                ),

            "files":
                len(combined)

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Could not process multiple files: {error}"

        }


# ============================================================
# STUDY TOOLS API
# ============================================================

@app.post("/api/study-tools")
async def study_tools(
    request: Request
):

    try:

        # ----------------------------------------------------
        # READ JSON MANUALLY
        #
        # This prevents FastAPI/Pydantic validation from
        # rejecting the request with HTTP 422 before the
        # endpoint can process it.
        # ----------------------------------------------------

        data = await request.json()

        if not isinstance(
            data,
            dict
        ):

            return {

                "success":
                    False,

                "error":
                    "Invalid Study Tools request."

            }

        # ----------------------------------------------------
        # ACTION
        # ----------------------------------------------------

        action = str(

            data.get(
                "action",
                ""
            )

            or ""

        ).lower().strip()

        # ----------------------------------------------------
        # CONTENT
        # ----------------------------------------------------

        content = str(

            data.get(
                "content",
                ""
            )

            or ""

        )

        # ----------------------------------------------------
        # TOPICS
        # ----------------------------------------------------

        topics = data.get(
            "topics",
            []
        )

        if not isinstance(
            topics,
            list
        ):

            topics = []

        topics = [

            str(topic).strip()

            for topic in topics

            if str(topic).strip()

        ]

        # ----------------------------------------------------
        # WEAK TOPICS
        # ----------------------------------------------------

        weak_topics = data.get(
            "weak_topics",
            []
        )

        if not isinstance(
            weak_topics,
            list
        ):

            weak_topics = []

        weak_topics = [

            str(topic).strip()

            for topic in weak_topics

            if str(topic).strip()

        ]

        # ----------------------------------------------------
        # OTHER VALUES
        # ----------------------------------------------------

        topic = str(

            data.get(
                "topic",
                ""
            )

            or ""

        ).strip()

        question = str(

            data.get(
                "question",
                ""
            )

            or ""

        ).strip()

        language = str(

            data.get(
                "language",
                "English"
            )

            or "English"

        ).strip()

        difficulty = str(

            data.get(
                "difficulty",
                "medium"
            )

            or "medium"

        ).strip()

        requested_format = str(

            data.get(
                "format",
                "exam"
            )

            or "exam"

        ).strip()

        # ----------------------------------------------------
        # SAFE NUMBER CONVERSION
        # ----------------------------------------------------

        try:

            days = int(
                data.get(
                    "days",
                    7
                )
            )

        except Exception:

            days = 7

        try:

            daily_minutes = int(
                data.get(
                    "daily_minutes",
                    45
                )
            )

        except Exception:

            daily_minutes = 45

        days = max(
            1,
            min(
                days,
                365
            )
        )

        daily_minutes = max(
            5,
            min(
                daily_minutes,
                1440
            )
        )

        # ----------------------------------------------------
        # CONTENT CHECK
        # ----------------------------------------------------

        if not content.strip():

            return {

                "success":
                    False,

                "error":
                    "Please provide study material first."

            }

        # ----------------------------------------------------
        # ALLOWED ACTIONS
        # ----------------------------------------------------

        allowed_actions = {

            "planner",
            "notes",
            "revision",
            "quiz",
            "tutor",
            "multilingual"

        }

        if action not in allowed_actions:

            return {

                "success":
                    False,

                "error":
                    f"Unknown study tool action: {action}"

            }

        # ----------------------------------------------------
        # CREATE PYDANTIC OBJECT AFTER NORMALIZATION
        # ----------------------------------------------------

        tool_request = StudyToolRequest(

            action=action,

            content=content,

            topics=topics,

            days=days,

            daily_minutes=daily_minutes,

            difficulty=difficulty,

            format=requested_format,

            topic=topic,

            weak_topics=weak_topics,

            question=question,

            language=language

        )

        # ----------------------------------------------------
        # GEMINI PROMPT
        # ----------------------------------------------------

        prompt = f"""

You are an expert AI learning assistant.

Task:

{action}

Student material:

{content[:50000]}

Topics:

{topics}

Difficulty:

{difficulty}

Requested format:

{requested_format}

Days:

{days}

Daily minutes:

{daily_minutes}

Revision topic:

{topic}

Weak topics:

{weak_topics}

Student question:

{question}

Target language:

{language}

Rules:

- Use only the supplied material for factual claims.
- Be student-friendly.
- Do not invent missing information.

For quiz:

Return JSON array with objects containing:

question

options

correct_answer

explanation

For other tasks:

Return concise useful text.

For multilingual output:

Faithfully translate the material's study content
into the requested language.

For planner/revision:

Make an actionable schedule.

"""

        # ----------------------------------------------------
        # QUIZ RESPONSE FORMAT
        # ----------------------------------------------------

        response_format = None

        if action == "quiz":

            response_format = {

                "type":
                    "json_schema",

                "json_schema": {

                    "name":
                        "quiz",

                    "schema": {

                        "type":
                            "array",

                        "items": {

                            "type":
                                "object",

                            "properties": {

                                "question": {

                                    "type":
                                        "string"

                                },

                                "options": {

                                    "type":
                                        "array",

                                    "items": {

                                        "type":
                                            "string"

                                    }

                                },

                                "correct_answer": {

                                    "type":
                                        "string"

                                },

                                "explanation": {

                                    "type":
                                        "string"

                                }

                            },

                            "required": [

                                "question",
                                "options",
                                "correct_answer",
                                "explanation"

                            ]

                        }

                    }

                }

            }

        # ----------------------------------------------------
        # GEMINI + FALLBACK
        # ----------------------------------------------------

        result = safe_gemini_call(

            prompt,

            response_format=response_format,

            fallback=study_tool_fallback(

                action,

                tool_request

            )

        )

        # ----------------------------------------------------
        # GEMINI FAILURE
        # ----------------------------------------------------

        if not result["success"]:

            return {

                "success":
                    False,

                "error":
                    result["error"]

            }

        # ----------------------------------------------------
        # PROCESS OUTPUT
        # ----------------------------------------------------

        output = result["output"]

        if (

            action == "quiz"

            and isinstance(
                output,
                str
            )

        ):

            try:

                output = json.loads(
                    output
                )

            except Exception:

                pass

        # ----------------------------------------------------
        # SUCCESS
        # ----------------------------------------------------

        return {

            "success":
                True,

            "output":
                output,

            "demo":
                result.get(
                    "demo",
                    False
                )

        }

    except Exception as error:

        return {

            "success":
                False,

            "error":
                f"Study tool failed: {str(error)}"

        }


# ============================================================
# RUN SERVER
# ============================================================

if __name__ == "__main__":

    import uvicorn

    uvicorn.run(

        app,

        host="127.0.0.1",

        port=8000

    )